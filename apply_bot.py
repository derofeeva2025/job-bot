#!/usr/bin/env python3
"""Listens for the "Apply" button under vacancy posts and prepares an application draft.

Runs on the owner's Mac (needs the `claude` CLI, logged in). For each press it
  1. reads the vacancy (link comes from the post's "View" button, text from the post + the page),
  2. asks `claude -p` to adapt ~/.claude/skills/apply/resumes/base.md, write a cover letter and answer any
     application questions found in the posting,
  3. sends the resume (.docx), cover letter (.docx) and the answers to the owner's private chat for review.
Nothing is ever sent to an employer: after "Approve" the owner submits it by hand.

Env (or a .env file next to this script): TELEGRAM_BOT_TOKEN, TELEGRAM_OWNER_ID (your numeric Telegram user id).
Run: python3 apply_bot.py        (keep it running; see apply_bot.plist for launchd)
Note: open a chat with the bot and press Start once, otherwise Telegram will not let it message you privately.
"""
import html, json, os, re, shutil, subprocess, sys, threading, time, urllib.error, urllib.request, uuid

BASE = os.path.dirname(os.path.abspath(__file__))
RESUME = os.path.expanduser("~/.claude/skills/apply/resumes/base.md")
OUT_DIR = os.path.expanduser("~/Documents/job-applications")
CLAUDE = shutil.which("claude") or os.path.expanduser("~/.local/bin/claude")
CLAUDE_TIMEOUT = 420

RULES = """You prepare a job application for Daria. Use ONLY facts from the base resume below.
- NEVER invent experience, skills, employers, dates, numbers or certifications. Reword and reorder only.
- Never name the employer JAKALA: write "International digital transformation agency" / "An international company" instead.
- Pick the angle (QA, Project Management or a mix) that matches the posting and emphasise that side.
- Resume: same facts, mirror the posting's vocabulary, most relevant experience first, tuned summary, same length (1-2 pages).
  Markdown: '# Name' title, '## SECTION' headings, '- ' bullets, **bold** for role lines.
- Language of resume and letter = language of the posting (German postings -> plain German, her German is B1, do not overstate it).
- Cover letter: SHORT, about 3 brief paragraphs, specific to this company and role, 2-3 concrete achievements tied to the main
  requirements, no clichés. The LAST line must say she is based in Europe (Switzerland) with stable power and internet (no outages),
  in the posting's language.
- If the posting contains application questions (or asks for salary expectations, notice period, work permit, motivation...), answer
  each briefly and honestly in the posting's language. If there are none, return an empty list.
- If a requirement is a true gap, do not fake it; put it in "gaps".
Return ONLY one JSON object, no prose, no code fences:
{"company": str, "role": str, "language": "en|de|...", "angle": "QA|PM|mix", "match_percent": int,
 "changes": [3-6 short strings in Russian describing what was changed in the resume],
 "gaps": [short strings in Russian],
 "resume_md": str, "cover_letter": str, "answers": [{"question": str, "answer": str}]}
"""

_env_file = os.path.join(BASE, ".env")
if os.path.exists(_env_file):
    for line in open(_env_file):
        if "=" in line and not line.lstrip().startswith("#"):
            k, _, v = line.strip().partition("=")
            os.environ.setdefault(k, v.strip().strip("\"'"))

TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN")
OWNER = os.environ.get("TELEGRAM_OWNER_ID")
JOBS = {}        # short id -> {"link", "text", "feedback", "dir"} for the Approve / Rewrite buttons
AWAIT_FIX = {}   # chat id -> job id, while waiting for the owner's rewrite instructions
LOCK = threading.Lock()  # one claude run at a time


def api(method, payload=None, files=None):
    url = f"https://api.telegram.org/bot{TOKEN}/{method}"
    if files:
        b = uuid.uuid4().hex
        body = b""
        for k, v in (payload or {}).items():
            body += f'--{b}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode()
        for k, (name, data) in files.items():
            body += (f'--{b}\r\nContent-Disposition: form-data; name="{k}"; filename="{name}"\r\n'
                     f'Content-Type: application/octet-stream\r\n\r\n').encode() + data + b"\r\n"
        body += f"--{b}--\r\n".encode()
        req = urllib.request.Request(url, data=body, headers={"Content-Type": f"multipart/form-data; boundary={b}"})
    else:
        req = urllib.request.Request(url, data=json.dumps(payload or {}).encode(), headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=90) as r:
            return json.loads(r.read())["result"]
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"Telegram {method} {e.code}: {e.read().decode(errors='replace')}") from None


def say(chat, text, markup=None):
    for i in range(0, len(text) or 1, 3900):  # Telegram limit is 4096
        p = {"chat_id": chat, "text": text[i:i + 3900], "disable_web_page_preview": True}
        if markup and i + 3900 >= len(text):
            p["reply_markup"] = markup
        api("sendMessage", p)


def page_text(url):
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 job-bot"})
        raw = urllib.request.urlopen(req, timeout=20).read(600_000).decode("utf-8", "replace")
    except Exception:
        return ""
    raw = re.sub(r"(?is)<(script|style|svg|nav|footer)[^>]*>.*?</\1>", " ", raw)
    t = re.sub(r"(?i)<br\s*/?>|</p>|</li>|</h\d>", "\n", raw)
    t = html.unescape(re.sub(r"<[^>]+>", " ", t))
    return re.sub(r"\n\s*\n+", "\n", re.sub(r"[ \t\xa0]+", " ", t)).strip()[:14000]


def run_claude(job_link, job_text, feedback=""):
    prompt = (f"{RULES}\n\n=== BASE RESUME ===\n{open(RESUME, encoding='utf-8').read()}\n\n"
              f"=== VACANCY LINK ===\n{job_link}\n\n=== VACANCY TEXT (from the Telegram post and the page; may be partial) ===\n{job_text}\n")
    if feedback:
        prompt += f"\n=== OWNER'S CHANGE REQUEST FOR THE PREVIOUS DRAFT (apply it) ===\n{feedback}\n"
    prompt += "\nIf the text above is thin, open the link with WebFetch to read the full posting. Output the JSON now."
    r = subprocess.run([CLAUDE, "-p", "--allowedTools", "WebFetch"], input=prompt, capture_output=True, text=True, timeout=CLAUDE_TIMEOUT)
    if r.returncode:
        raise RuntimeError(f"claude failed: {r.stderr.strip()[:500]}")
    s = r.stdout
    try:
        return json.loads(s[s.index("{"):s.rindex("}") + 1])
    except ValueError:
        raise RuntimeError("claude did not return valid JSON:\n" + s[:500]) from None


def md_to_docx(md, path, title_size=18):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    for line in md.splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(); r = p.add_run(line[2:].strip()); r.bold = True; r.font.size = Pt(title_size)
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(); r = p.add_run(line[3:].strip().upper()); r.bold = True; r.font.size = Pt(12)
            continue
        bullet = line.lstrip().startswith(("- ", "* "))
        p = doc.add_paragraph(style="List Bullet" if bullet else None)
        for i, part in enumerate(re.split(r"\*\*", line.lstrip()[2:] if bullet else line)):
            p.add_run(part).bold = bool(i % 2)
    doc.save(path)


def safe(s):
    return re.sub(r"[^\w.-]+", "_", s).strip("_")[:40] or "x"


def process(chat, jid):
    job = JOBS[jid]
    with LOCK:
        say(chat, "⏳ Готовлю черновик (1–3 минуты)…")
        text = job["text"] + "\n\n" + page_text(job["link"])
        d = run_claude(job["link"], text[:20000], job.get("feedback", ""))
    folder = os.path.join(OUT_DIR, f"{safe(d.get('company', ''))}_{safe(d.get('role', ''))}")
    os.makedirs(folder, exist_ok=True)
    job["dir"] = folder
    resume_p, letter_p = os.path.join(folder, "resume.docx"), os.path.join(folder, "cover_letter.docx")
    md_to_docx(d["resume_md"], resume_p)
    md_to_docx(d["cover_letter"], letter_p)
    for n, c in (("resume.md", d["resume_md"]), ("cover_letter.md", d["cover_letter"]), ("answers.json", json.dumps(d.get("answers", []), ensure_ascii=False, indent=1))):
        open(os.path.join(folder, n), "w", encoding="utf-8").write(c)
    label = f"{d.get('company', '?')} — {d.get('role', '?')}"
    for p, cap in ((resume_p, f"Резюме: {label}"), (letter_p, "Cover letter")):
        api("sendDocument", {"chat_id": chat, "caption": cap}, {"document": (f"{os.path.basename(folder)}_{os.path.basename(p)}", open(p, "rb").read())})
    msg = [f"📋 {label}", f"Ракурс: {d.get('angle', '')}  |  Совпадение: {d.get('match_percent', '?')}%", ""]
    if d.get("changes"):
        msg += ["Что изменено в резюме:"] + [f"• {c}" for c in d["changes"]] + [""]
    if d.get("gaps"):
        msg += ["Чего не хватает:"] + [f"• {g}" for g in d["gaps"]] + [""]
    msg += ["✉️ Cover letter:", d["cover_letter"].strip()]
    if d.get("answers"):
        msg += ["", "❓ Ответы на вопросы вакансии:"]
        for a in d["answers"]:
            msg += [f"— {a['question']}", a["answer"], ""]
    msg += ["", f"Файлы: {folder}"]
    say(chat, "\n".join(msg), {"inline_keyboard": [[
        {"text": "✅ Одобрить", "callback_data": f"ok:{jid}"},
        {"text": "🔁 Переписать", "callback_data": f"fix:{jid}"},
        {"text": "✖️ Отмена", "callback_data": f"no:{jid}"}]]})


def safe_process(chat, jid):
    try:
        process(chat, jid)
    except Exception as e:
        print("apply failed:", e)
        try:
            say(chat, f"⚠️ Не получилось подготовить черновик: {e}"[:1500])
        except Exception:
            pass


def on_callback(cq):
    uid, data, msg = str(cq["from"]["id"]), cq.get("data", ""), cq.get("message") or {}
    if not OWNER or uid != OWNER:
        print(f"callback from user id {uid} ignored (set TELEGRAM_OWNER_ID={uid} if this is you)")
        api("answerCallbackQuery", {"callback_query_id": cq["id"], "text": "Эта кнопка только для владельца бота.", "show_alert": True})
        return
    if data == "apply":
        link = next((b["url"] for row in (msg.get("reply_markup") or {}).get("inline_keyboard", []) for b in row if "url" in b and "View" in b["text"]), "")
        jid = uuid.uuid4().hex[:10]
        JOBS[jid] = {"link": link, "text": msg.get("text", "")}
        try:
            api("answerCallbackQuery", {"callback_query_id": cq["id"], "text": "Готовлю черновик — пришлю в личку бота."})
            api("sendChatAction", {"chat_id": uid, "action": "typing"})  # fails if the bot was never started
        except RuntimeError as e:
            print(e)
            return
        threading.Thread(target=safe_process, args=(uid, jid), daemon=True).start()
        return
    kind, _, jid = data.partition(":")
    job = JOBS.get(jid)
    api("answerCallbackQuery", {"callback_query_id": cq["id"]})
    if not job:
        say(uid, "Черновик устарел (бот перезапускался). Нажмите Apply под вакансией ещё раз.")
    elif kind == "ok":
        say(uid, f"✅ Готово к отправке. Откройте вакансию и подайте заявку вручную:\n{job['link']}\n\nФайлы: {job.get('dir', '')}")
    elif kind == "fix":
        AWAIT_FIX[uid] = jid
        say(uid, "Напишите одним сообщением, что изменить (например: «письмо короче», «сделай упор на PM»).")
    elif kind == "no":
        JOBS.pop(jid, None)
        say(uid, "Отменено.")


def on_message(m):
    uid = str(m["from"]["id"])
    if uid != OWNER:
        return
    if m.get("text", "").startswith("/start"):
        say(uid, "Привет! Нажимайте ✍️ Apply под вакансией в канале — черновик придёт сюда.")
    elif uid in AWAIT_FIX and m.get("text"):
        jid = AWAIT_FIX.pop(uid)
        if jid in JOBS:
            JOBS[jid]["feedback"] = m["text"]
            threading.Thread(target=safe_process, args=(uid, jid), daemon=True).start()


def main():
    if not TOKEN:
        sys.exit("TELEGRAM_BOT_TOKEN not set")
    if not OWNER:
        print("TELEGRAM_OWNER_ID is not set: press Apply once, the log will show your id; then put it in .env")
    if not os.path.exists(RESUME):
        sys.exit(f"base resume not found: {RESUME}")
    api("deleteWebhook")  # getUpdates does not work while a webhook is set
    offset = None
    print("apply bot is listening…")
    while True:
        try:
            ups = api("getUpdates", {"timeout": 50, "offset": offset, "allowed_updates": ["callback_query", "message"]})
        except Exception as e:
            print("poll error:", e)
            time.sleep(5)
            continue
        for u in ups:
            offset = u["update_id"] + 1
            try:
                if "callback_query" in u:
                    on_callback(u["callback_query"])
                elif "message" in u:
                    on_message(u["message"])
            except Exception as e:
                print("handler error:", e)


if __name__ == "__main__":
    main()
